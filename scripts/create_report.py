"""Build the local course report from its retained template and actual evidence.

Usage: python -X utf8 scripts/create_report.py
Re-running replaces only the selected output DOCX. No input is modified or deleted.
PDF export and page rendering are separate steps, see scripts/export_report.ps1.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
TEST_LOGS = [
    ("基础 TCP", "test_basic.txt"),
    ("聊天室协议", "test_chat.txt"),
    ("直接文字会话", "directcall_test.txt"),
    ("服务器对象", "chatserver_test.txt"),
    ("Qt 界面场景", "ui_test.txt"),
]


def read_text(path: Path) -> str:
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    return data.decode("utf-8-sig")


def optional_json(path: Path) -> dict:
    if not path.is_file():
        return {}
    value = json.loads(read_text(path))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path.name}")
    return value


def test_result(path: Path) -> dict:
    """Count observable test cases, excluding Qt suite setup/cleanup hooks."""
    if not path.is_file():
        return {"passed": 0, "failed": 0, "skipped": 0, "status": "缺失待验证"}
    text = read_text(path)
    if "Start testing of" in text:
        cases = re.findall(r"^(PASS|FAIL!|SKIP)\s*:\s*([^\r\n]+)", text, re.M)
        cases = [(status, name) for status, name in cases
                 if not re.search(r"::(?:initTestCase|cleanupTestCase)\(\)", name)]
        passed = sum(status == "PASS" for status, _ in cases)
        failed = sum(status == "FAIL!" for status, _ in cases)
        skipped = sum(status == "SKIP" for status, _ in cases)
        totals = re.search(r"Totals:\s*(\d+) passed,\s*(\d+) failed,\s*(\d+) skipped", text)
        complete = totals is not None and int(totals[2]) == 0 and "Finished testing" in text
        ok = complete and failed == 0 and skipped == 0 and passed > 0
    else:
        totals = re.search(r"^Ran (\d+) tests? in [^\r\n]+", text, re.M)
        count = int(totals[1]) if totals else 0
        failure = re.search(r"^FAILED \(([^)]+)\)", text, re.M)
        stats = dict(re.findall(r"(failures|errors|skipped)=(\d+)", failure[1] if failure else ""))
        failed = int(stats.get("failures", 0)) + int(stats.get("errors", 0))
        skipped_match = re.search(r"^OK \(skipped=(\d+)\)", text, re.M)
        skipped = int(stats.get("skipped", skipped_match[1] if skipped_match else 0))
        passed = max(0, count - failed - skipped)
        ok = count > 0 and bool(re.search(r"^OK\s*$", text, re.M)) and not failed and not skipped
    return {"passed": passed, "failed": failed, "skipped": skipped,
            "status": "通过" if ok else "存在失败或未完整结束"}


def font(run, name="宋体", size=11, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def add_styles(doc):
    for name, face, size, bold in [
        ("ReportBody", "宋体", 11, False),
        ("ReportHeading", "黑体", 14, True),
        ("ReportSubheading", "黑体", 12, True),
        ("ReportCaption", "宋体", 10, False),
    ]:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = face
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), face)
        pf = style.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.25
        pf.first_line_indent = Cm(0)
        pf.keep_with_next = name in ("ReportHeading", "ReportSubheading")
        pf.widow_control = True
        # The retained template has a Chinese line grid. Body paragraphs use
        # their explicit line spacing instead of snapping to that cover grid.
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        style.element.get_or_add_pPr().append(snap)
    doc.styles["ReportHeading"].paragraph_format.space_after = Pt(10)


class Report:
    def __init__(self, template: Path, evidence: Path):
        self.doc = Document(template)
        self.evidence = evidence
        self.results = [(label, name, test_result(evidence / name)) for label, name in TEST_LOGS]
        self.runtime = optional_json(evidence / "runtime_verification.json")
        self.capture = optional_json(evidence / "capture_analysis.json")
        self.analysis_text = read_text(evidence / "tcp_analysis.md") if (evidence / "tcp_analysis.md").is_file() else ""
        self.basic_text = read_text(evidence / "basic_conversation.txt") if (evidence / "basic_conversation.txt").is_file() else ""
        self.missing = []
        self.next_heading_starts_page = False
        self.prepare_template()

    def prepare_template(self):
        paragraphs = self.doc.paragraphs
        body_start = next((p for p in paragraphs if p.text.strip() == "实验内容"), None)
        if body_start is None:
            raise ValueError("Template must contain the original 实验内容 section")
        # Keep the existing cover and section properties; replace only body slots.
        body = self.doc._element.body
        started = False
        for child in list(body):
            started = started or child is body_start._p
            if started and child.tag != qn("w:sectPr"):
                body.remove(child)
        for p in self.doc.paragraphs:
            for label, value in [("实验名称：", "Socket通信编程实验"), ("课程名称：", "计算机网络"),
                                 ("姓名：", "待填写"), ("学号：", "待填写"), ("日期：", "2026-10-13")]:
                if p.text.startswith(label):
                    # Preserve the label and the template's underline run formatting.
                    p.runs[0].text = label
                    if len(p.runs) < 2:
                        p.add_run()
                    p.runs[1].text = value
                    for run in p.runs[2:]:
                        run.text = ""
                    if label == "实验名称：":
                        p.runs[1].font.size = Pt(16)
                    break
        add_styles(self.doc)
        for section in self.doc.sections:
            borders = section._sectPr.find(qn("w:pgBorders"))
            if borders is not None:
                # Keep the template's no-border cover and black body border,
                # but separate all four rules from adjacent text strokes.
                borders.set(qn("w:offsetFrom"), "text")
                for edge in borders:
                    edge.set(qn("w:space"), "8")
        self.doc.core_properties.title = "实验6 Socket通信编程实验报告"
        self.doc.core_properties.subject = "计算机网络课程实验"
        self.doc.core_properties.author = "待填写"
        self.doc.core_properties.last_modified_by = "待填写"
        self.doc.core_properties.comments = ""

    def heading(self, text, sub=False):
        p = self.doc.add_paragraph(text, "ReportSubheading" if sub else "ReportHeading")
        if self.next_heading_starts_page:
            p.paragraph_format.page_break_before = True
            self.next_heading_starts_page = False

    def paragraph(self, text, indent=True):
        p = self.doc.add_paragraph(text, "ReportBody")
        if indent:
            p.paragraph_format.first_line_indent = Cm(0.75)
        return p

    def page(self):
        # An explicit break in a new empty paragraph can itself overflow a full
        # page, producing a blank page. Put the break on the next real heading.
        self.next_heading_starts_page = True

    def table(self, headers, rows, widths):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = str(text)
        for row in rows:
            for cell, text in zip(table.add_row().cells, row):
                cell.text = str(text)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + edge)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "D9D9D9")
            borders.append(el)
        table._tbl.tblPr.append(borders)
        for i, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if i == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if i == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E7EEF5")
                    cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.style = self.doc.styles["ReportBody"]
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.line_spacing = 1.1
                    for run in p.runs:
                        font(run, size=10.5, bold=i == 0)

    def image(self, name, caption, width=15.2, max_height=10.7):
        path = self.evidence / name
        if not path.is_file():
            self.missing.append(name)
            self.paragraph("待补充截图：" + name, False)
            return
        with Image.open(path) as im:
            width = min(width, max_height * im.width / im.height)
        p = self.doc.add_paragraph(style="ReportBody")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width))
        p = self.doc.add_paragraph(caption, "ReportCaption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def capture_item(self, name):
        return next((c for c in self.capture.get("captures", []) if c.get("file") == name), {})

    def build(self):
        # Page 2: keep the template's first two chapter intentions.
        self.heading("一 实验内容")
        self.paragraph("本实验完成了基础 TCP 一对一通信和 Qt 图形聊天室。基础程序验证 socket 创建、绑定、监听、连接及双向收发；聊天室在持续 TCP 连接上实现昵称登记、在线成员、群组广播、私聊和文件发送，并通过服务器呼叫信令建立客户端之间的独立 TCP 文字会话。")
        self.table(["部分", "实现内容", "通信路径"], [
            ["基础 TCP", "客户端先发、服务器回复；/quit 结束", "客户端 ↔ 服务器"],
            ["聊天室", "公共大厅、群组、私聊、文件", "客户端 ↔ 聊天服务器"],
            ["直接文字", "请求、接听、关联、文字、挂断", "信令经服务器；文字客户端直连"],
        ], [2.4, 8.0, 6.6])
        self.paragraph("所有应用帧以四字节大端长度字段标记边界。基础负载为文字，Qt 负载为 UTF-8 JSON 对象。项目不包含账号密码、数据库历史、语音视频或 NAT 穿透；群组与成员状态仅存服务器内存。")
        self.heading("二 实验环境")
        self.table(["项目", "本次环境与参数"], [
            ["操作系统", "Windows x64，同一台电脑回环地址 127.0.0.1"],
            ["开发工具", "Qt 5.15.2，MSVC 19.50，C++17；qmake 与 nmake"],
            ["Qt 模块", "Core、Gui、Widgets、Network；测试使用 QtTest"],
            ["端口", "默认基础 8888、聊天室 8889；抓包用 18888/18889 和临时直连端口"],
            ["容量与超时", "文件 ≤ 5 MiB；聊天室与直连帧 ≤ 10 MiB；响铃 30 秒，直连关联 10 秒"],
            ["辅助工具", "Python 网络测试、Wireshark/Npcap 抓包、Word 模板报告"],
        ], [3.1, 13.9])
        self.paragraph("QtTest 所列 windows 10 为运行库报告的内核信息，未据此指定系统发行名称。验证范围为同机 loopback；双机局域网和另一台全新 Windows 系统尚未实测。")

        # Page 3: basic-program implementation, actual conversation, server overview.
        self.page()
        self.heading("三 实验步骤")
        self.heading("基础 TCP 的建立 收发和关闭", True)
        self.paragraph("服务器调用 socket、bind、listen、accept；客户端调用 socket、connect。收发循环先处理四字节长度，再按长度读写全部正文，处理部分 send/recv、对端关闭和截断帧。输入 /quit 或终端输入结束时释放连接。")
        basic = self.capture_item("basic_session.pcapng")
        stream = next(iter(basic.get("streams", [])), {})
        messages = stream.get("messages", [])
        requests = [m["body"] for m in messages if m.get("direction") == 0 and isinstance(m.get("body"), str) and m["body"] != "/quit"]
        replies = [m["body"] for m in messages if m.get("direction") == 1 and isinstance(m.get("body"), str)]
        if self.basic_text and requests and replies:
            self.paragraph("真实基础程序在 18888 端口完成以下两轮交流，标准输出保存在 basic_conversation.txt，文字负载同时由 basic_session.pcapng 重组得到：")
            self.table(["轮次", "客户端发送", "服务器回复"],
                       [[i + 1, a, b] for i, (a, b) in enumerate(zip(requests[:2], replies[:2]))], [1.2, 7.8, 8.0])
            self.paragraph("随后客户端发送 /quit，服务器返回关闭提示，双方正常结束。该过程既验证收发方向，也验证中文文字和消息边界。")
        else:
            self.paragraph("基础双向交流记录尚不完整，待补充 basic_conversation.txt 及实际抓包重组结果。")
        self.heading("构建和启动聊天室", True)
        self.paragraph("执行 run_build.ps1 构建基础程序、Qt 主程序及测试。启动一个服务器和三个不同昵称的客户端；昵称登记成功后显示在线列表。服务器负责成员、群组与呼叫状态；界面按连接和登录状态控制操作。")
        self.image("06_server.png", "图1 图形服务器记录用户 群组和呼叫事件", width=11.0, max_height=7.7)

        # Page 4: two full GUI views.
        self.page()
        self.heading("四 实验结果与分析")
        self.heading("公共大厅 私聊与群组隔离", True)
        self.paragraph("Alice、Bob、Charlie 登录后均可收到大厅消息。Alice 对 Bob 的私聊只出现在双方窗口；新建讨论组由 Alice 和 Bob 加入，Charlie 不接收讨论组文字。界面测试同时检查消息可见性。")
        self.image("01_lobby.png", "图2 Charlie 的大厅视图  不出现 Alice 发给 Bob 的私聊", width=12.0, max_height=8.4)
        self.image("02_groups.png", "图3 Alice 的实验讨论组  成员为 Alice 和 Bob", width=12.0, max_height=8.4)

        # Page 5: file outcome, log-derived regression counts, isolated runtime.
        self.page()
        self.heading("文件保存与回归验证", True)
        self.image("03_files.png", "图4 Bob 收到群内文件  中文二进制文件与空文件均保存", width=14.0, max_height=9.8)
        self.paragraph("UI 场景对中文文件名的二进制内容逐字节比较，并确认空文件大小为零。文件以 Base64 进入 JSON，服务器按原始字节限制 5 MiB。发送者的转发提示与接收者的保存结果分别对应网络转发和文件落盘。")
        rows = [[label, r["passed"], r["failed"], r["skipped"], r["status"]]
                for label, _, r in self.results]
        self.table(["测试范围", "通过", "失败", "跳过", "结果"], rows, [4.3, 1.3, 1.3, 1.3, 8.8])
        passed = sum(r["passed"] for _, _, r in self.results)
        all_ok = all(r["status"] == "通过" for _, _, r in self.results)
        self.paragraph(f"以上计数直接读取五份证据日志，Qt 的初始化与清理项不计入行为测试。本次共 {passed} 个行为测试通过。"
                       + ("各套日志完整结束且未发现失败或跳过。" if all_ok else "仍有缺失或失败的测试记录，不能据此宣布整套通过。"))
        runtime_ok = self.runtime.get("status") == "PASS" and len(self.runtime.get("gui_processes", [])) == 4
        if runtime_ok:
            self.paragraph("运行包在去掉 Qt/conda 开发 PATH 的环境中通过基础和聊天室测试，并启动一个图形服务器与三个独立图形客户端。runtime_verification.json 记录四个进程的 Qt DLL 均来自运行包；该检查仍在同一台电脑上完成。")
        else:
            self.paragraph("脱离开发 PATH 的四个独立 GUI 进程验证尚无完整 PASS，运行包独立运行结论待验证。")

        # Page 6: actual dialogs and direct transport.
        self.page()
        self.heading("呼叫与客户端直接 TCP 文字会话", True)
        self.paragraph("Alice 选择 Bob 发起呼叫，服务器向 Bob 发送来电信令。Bob 接听后先在 IPv4 临时端口监听，再把端口返回服务器；Alice 连接该端口，hello/ready 按 callId 关联成功后，双方才打开直接文字窗口。")
        self.image("04_incoming_call.png", "图5 Bob 的来电提示  支持接听或拒绝", width=9.0, max_height=3.2)
        self.image("05_direct_call.png", "图6 Alice 的独立 TCP 文字窗口  双向中文收发", width=13.0, max_height=9.0)
        self.paragraph("本次双方各发一句中文，并在挂断后继续群聊。测试还覆盖拒接、取消、忙线/离线信令、未接听超时、拒绝直连、关联超时、错误 callId、拆包粘包和会话复用。响铃由服务器限时 30 秒，接听后的 TCP 建连与关联由 DirectCall 限时 10 秒。")
        self.paragraph("直接会话与私聊使用不同路径：私聊始终由原聊天室 TCP 连接经服务器转发，接听后的独立文字由新 TCP 连接发送。下一页按实际抓包端点复核这一差别。")

        # Page 7: evidence-driven packet analysis with independent stream numbers.
        self.page()
        self.heading("TCP 抓包与消息边界分析", True)
        captures = self.capture.get("captures", [])
        capture_ok = self.capture.get("status") == "PASS" and bool(captures) and all((self.evidence / c["file"]).is_file() for c in captures)
        if capture_ok:
            summary = self.capture.get("summary", {})
            self.paragraph(f"本次只捕获实验端口的 Npcap loopback 流量。{len(captures)} 份捕获中共分析 {summary.get('tcp_streams', '待确认')} 条 TCP 流，"
                           f"观察到 {summary.get('handshakes', '待确认')} 次握手，{summary.get('fin_closed_streams', '待确认')} 条流出现 FIN 关闭；各文件的流编号独立。")
            self.table(["捕获文件", "TCP包", "流数"],
                       [[c["file"], c.get("packet_count", "待确认"), len(c.get("streams", []))] for c in captures], [11.3, 2.7, 3.0])
            selected = []
            for name in ("basic_session.pcapng", "direct_call.pcapng"):
                for s in self.capture_item(name).get("streams", []):
                    selected.append(["基础" if name.startswith("basic") else "直连", f"{s.get('syn_frames', [])} / {s.get('syn_ack_frames', [])} / {s.get('first_ack_frame', '待确认')}",
                                     str(s.get("fin_frames", [])), str(s.get("rst_frames", []))])
            self.table(["流类型", "SYN / SYN ACK / ACK帧", "FIN帧", "RST帧"], selected, [2.4, 8.3, 4.0, 2.3])
            for name in ("basic_session.pcapng", "direct_call.pcapng"):
                for s in self.capture_item(name).get("streams", []):
                    self.paragraph(f"{name} 的流 {s.get('stream')}：{s.get('node0')} ↔ {s.get('node1')}。", False)
            if self.analysis_text:
                lines = [line.strip() for line in self.analysis_text.splitlines() if line.startswith("基础程序使用")]
                if lines:
                    self.paragraph(lines[0])
            else:
                self.paragraph("tcp_analysis.md 尚未提供；详细逐包解释待补充。")
            types = summary.get("direct_call_payload_types", [])
            self.paragraph("直接流的重组消息类型为 " + "、".join(types) + "。其端点使用被叫临时端口 "
                           + str(summary.get("direct_call_port", "待确认")) + "，与聊天室监听端口不同。")
            if summary.get("server_payload_has_call_chat") is False and "call_chat" in types:
                self.paragraph("聊天室服务器流中没有 call_chat，独立直连流中存在双向 call_chat，两者共同支持文字经客户端直接 TCP 通道传输。ACK 可与 FIN 等标志合并，关闭过程按实际帧说明，不强行套成四个独立包。")
            self.paragraph("复核时在 Wireshark 中用 tcp.stream eq N 筛选，查看端口、Flags、Seq/Ack，再追踪 TCP 流读取长度字段与正文。完整帧号、重组和关闭解释见 tcp_analysis.md；当前捕获为明文，未启用 TLS。")
        else:
            self.paragraph("抓包分析或原始 pcapng 尚不完整，当前仅列出复核方法，不声称已完成握手、关闭或直接通道验证。待补充 capture_analysis.json、tcp_analysis.md 和对应 pcapng。")
            self.paragraph("复核方法为：检查三次握手和实际 FIN/RST 关闭，重组 TCP 字节流后读四字节大端长度，比较服务器信令流和直接文字流的端点与消息类型。")

        # Page 8: conclusions, limitations, authentic human team fields.
        self.page()
        self.heading("五 实验总结")
        self.paragraph("基础 TCP 实验将 socket 生命周期与应用层消息边界联系起来。分帧不能依赖一次 recv 的返回长度；完整收发循环必须同时考虑部分数据、正常关闭与截断。UTF-8 中文交流和拆包测试进一步验证了按字节长度而非字符数量处理消息的重要性。")
        self.paragraph("Qt 部分以事件驱动处理多客户端连接。把服务器业务、直接文字会话和界面状态分开后，群组成员检查决定转发范围，明确的登录回执决定客户端可发送状态，呼叫编号把服务器信令和直接握手关联起来。挂断后关闭资源并解除旧连接回调，保证后续会话可复用。")
        self.paragraph("文件收发分别检查网络转发和保存结果，避免把已发送误写为已保存。直接文字设计为通过服务器之外的新 TCP 流传输。"
                       + ("本次抓包中的端口、hello/ready 和双向文字提供了可复核的连接依据。" if capture_ok
                          else "实际连接端点与传输过程仍需完整抓包证据复核。"))
        self.heading("实际限制与后续验证", True)
        self.paragraph("本次验证范围为同一台 Windows x64 电脑的 loopback 环境。双机局域网需要额外确认防火墙和被叫临时端口可达；外网 NAT、语音视频、账号认证、加密和持久化历史不属于当前实现。"
                       + ("当前运行包已通过免开发 PATH 检查，另一台全新 Windows 电脑仍需实际试运行。" if runtime_ok
                          else "运行包的免开发 PATH 检查和另一台全新 Windows 电脑试运行仍待验证。"))
        self.heading("真实队内分工", True)
        self.paragraph("以下预留四行，由实际参与实验的同学按真实人数调整填写；姓名、学号和分工均保留待填写，不按程序模块或自动化执行过程推定人员贡献。")
        self.table(["姓名", "学号", "实际工作及核对结果"], [["待填写", "待填写", "待填写"] for _ in range(4)], [3.4, 4.5, 9.1])
        self.heading("复现与提交", True)
        self.paragraph("按 README 执行 run_build.ps1、run_tests.ps1 和 run_package.ps1；scripts/verify_runtime.py 验证运行包。run_capture.ps1 保存实验流量和分析，再运行 scripts/create_report.py 重建本报告，scripts/export_report.ps1 导出 PDF。提交前补齐个人资料并检查全部报告页。")
        self.paragraph("报告从现有测试日志、GUI 截图、基础交流记录、运行包检查和抓包分析中读取证据，缺失项在对应结果中说明。源代码与脱敏文档可公开，个人信息报告与原始证据留在本地交付目录。")
        return self.doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=ROOT / "实验报告模板.docx")
    parser.add_argument("--evidence", type=Path, default=ROOT / "delivery/evidence")
    parser.add_argument("--output", type=Path, default=ROOT / "delivery/report/实验6_Socket通信编程实验报告.docx")
    args = parser.parse_args()
    if args.template.resolve() == args.output.resolve():
        parser.error("Output must differ from the retained template")
    report = Report(args.template, args.evidence)
    doc = report.build()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"DOCX: {args.output}")
    for label, name, result in report.results:
        print(f"{name}: {result['passed']} behavior cases passed; {result['failed']} failed; {result['skipped']} skipped; {result['status']}")
    if report.missing:
        print("Missing screenshots: " + ", ".join(report.missing))
    print("PDF export and all-page visual inspection are required before submission.")


if __name__ == "__main__":
    main()
